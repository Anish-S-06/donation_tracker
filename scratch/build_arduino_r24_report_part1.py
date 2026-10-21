import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from build_arduino_r24_report import (
    set_cell_background,
    set_cell_margins,
    set_table_borders,
    set_callout_borders,
    make_row_cant_split,
    make_row_header,
    add_heading_1,
    add_heading_2,
    add_heading_3,
    add_body,
    add_bullet,
    add_callout,
    add_code_block,
    add_styled_table
)

def add_page_number_to_run(run):
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def build_document():
    doc = docx.Document()
    
    # Page Setup
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    section.different_first_page_header_footer = True
    
    # Running Header
    header = section.header
    p_head = header.paragraphs[0]
    p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_head = p_head.add_run("ARDUINO-BASED MULTI-SENSOR ENVIRONMENTAL MONITORING SYSTEM")
    r_head.font.name = "Times New Roman"
    r_head.font.size = Pt(8.5)
    r_head.font.color.rgb = RGBColor(128, 128, 128)
    
    # Running Footer
    footer = section.footer
    p_foot = footer.paragraphs[0]
    p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_foot_left = p_foot.add_run("Technical Capstone Project Report                                                                             Page ")
    r_foot_left.font.name = "Times New Roman"
    r_foot_left.font.size = Pt(9)
    r_foot_left.font.color.rgb = RGBColor(100, 116, 139)
    
    r_foot_pg = p_foot.add_run()
    r_foot_pg.font.name = "Times New Roman"
    r_foot_pg.font.size = Pt(9)
    r_foot_pg.font.bold = True
    r_foot_pg.font.color.rgb = RGBColor(30, 41, 59)
    add_page_number_to_run(r_foot_pg)

    # =========================================================================
    # FRONT MATTER: 1. TITLE PAGE
    # =========================================================================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("A TECHNICAL CAPSTONE PROJECT REPORT")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(71, 85, 105)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("ARDUINO-BASED MULTI-SENSOR\nENVIRONMENTAL MONITORING SYSTEM")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(36)
    r = p.add_run("Real-Time Microclimate Telemetry, Chemiresistive Gas Response Mapping, and Embedded Visual Interface")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.italic = True
    r.font.color.rgb = RGBColor(51, 65, 85)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Submitted in Partial Fulfilment of the Academic Requirements\nfor the Award of the Degree of")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(71, 85, 105)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(36)
    r = p.add_run("BACHELOR OF TECHNOLOGY / DIPLOMA IN ENGINEERING\nIN ELECTRONICS AND INSTRUMENTATION ENGINEERING / COMPUTER SCIENCE")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)

    # 2-column table for candidates and guides
    info_table = doc.add_table(rows=1, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False
    cell_l = info_table.cell(0, 0)
    cell_r = info_table.cell(0, 1)
    cell_l.width = Inches(3.25)
    cell_r.width = Inches(3.25)
    set_cell_margins(cell_l, top=60, bottom=60, left=40, right=40)
    set_cell_margins(cell_r, top=60, bottom=60, left=40, right=40)
    
    pl = cell_l.paragraphs[0]
    pl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    rl_title = pl.add_run("Submitted by:\n")
    rl_title.font.bold = True
    rl_title.font.size = Pt(11)
    rl_body = pl.add_run("[Student Name]\nRoll / Register No: [Roll Number]\nDepartment of [Department Name]")
    rl_body.font.size = Pt(11)
    
    pr = cell_r.paragraphs[0]
    pr.alignment = WD_ALIGN_PARAGRAPH.LEFT
    rr_title = pr.add_run("Under the Guidance of:\n")
    rr_title.font.bold = True
    rr_title.font.size = Pt(11)
    rr_body = pr.add_run("[Guide / Faculty Name]\n[Faculty Designation / Title]\nDepartment of [Department Name]")
    rr_body.font.size = Pt(11)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)

    # Institutional details table
    inst_table = doc.add_table(rows=1, cols=2)
    inst_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    inst_table.autofit = False
    c_il = inst_table.cell(0, 0)
    c_ir = inst_table.cell(0, 1)
    c_il.width = Inches(3.25)
    c_ir.width = Inches(3.25)
    set_cell_margins(c_il, top=60, bottom=60, left=40, right=40)
    set_cell_margins(c_ir, top=60, bottom=60, left=40, right=40)
    
    pil = c_il.paragraphs[0]
    pil.alignment = WD_ALIGN_PARAGRAPH.LEFT
    ril_title = pil.add_run("Institution Details:\n")
    ril_title.font.bold = True
    ril_title.font.size = Pt(11)
    ril_body = pil.add_run("[College / University Name]\n[Campus Location / City, State]")
    ril_body.font.size = Pt(11)
    
    pir = c_ir.paragraphs[0]
    pir.alignment = WD_ALIGN_PARAGRAPH.LEFT
    rir_title = pir.add_run("Academic Period:\n")
    rir_title.font.bold = True
    rir_title.font.size = Pt(11)
    rir_body = pir.add_run("Academic Session: [Academic Year]\nSubmission Date: October 2026")
    rir_body.font.size = Pt(11)

    # =========================================================================
    # FRONT MATTER: 2. CERTIFICATE & STUDENT DECLARATION
    # =========================================================================
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("CERTIFICATE OF RECOMMENDATION")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)

    add_body(doc, "This is to certify that the comprehensive technical capstone project report entitled **“Arduino-Based Multi-Sensor Environmental Monitoring System: Real-Time Microclimate Telemetry, Chemiresistive Gas Response Mapping, and Embedded Visual Interface”** represents the bona fide engineering investigation executed by **[Student Name]**, Roll Number **[Roll Number]**, under my direct supervision and technical guidance.")
    add_body(doc, "The project documents the complete architectural lifecycle of a localized environmental observation instrument, including virtual Electronic Design Automation (EDA) simulation within Autodesk Tinkercad Circuits, low-level AVR C++ firmware development, conversion mathematics, threshold mapping logic, physical breadboard fabrication, and benchtop experimental testing utilizing the Arduino Uno R3, DHT11 temperature/humidity sensor, MQ-135 metal oxide gas sensor, and a 16x2 I2C alphanumeric liquid crystal display.")
    add_body(doc, "The technical contents and experimental findings compiled within this document have not been submitted previously to any other university, board, or examining committee for the conferral of any academic degree, diploma, or institutional award.")

    # 3-column Signatures
    sig_table = doc.add_table(rows=1, cols=3)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False
    for i in range(3):
        sig_table.cell(0, i).width = Inches(2.16)
        set_cell_margins(sig_table.cell(0, i), top=120, bottom=60, left=40, right=40)
    
    p1 = sig_table.cell(0, 0).paragraphs[0]
    p1.add_run("_____________________\n").font.bold = True
    p1.add_run("[Guide Name]\nProject Supervisor\nDate: ___________")
    
    p2 = sig_table.cell(0, 1).paragraphs[0]
    p2.add_run("_____________________\n").font.bold = True
    p2.add_run("[HOD Name]\nHead of Department\nDate: ___________")
    
    p3 = sig_table.cell(0, 2).paragraphs[0]
    p3.add_run("_____________________\n").font.bold = True
    p3.add_run("External Examiner\nInstitutional Committee\nDate: ___________")

    # Student Declaration
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("STUDENT DECLARATION")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)

    add_body(doc, "I hereby declare that this project report titled **“Arduino-Based Multi-Sensor Environmental Monitoring System”** is an authentic record of original engineering work conducted by me. All circuit schematics, firmware algorithms, hardware configurations, and experimental observations presented herein reflect genuine investigatory efforts.")
    add_body(doc, "I explicitly affirm that all virtual simulation trials were performed within Autodesk Tinkercad Circuits and are clearly delineated from physical breadboard tests. The physical prototype utilizes an Arduino Uno, digital DHT11 composite sensor, MQ-135 chemiresistive gas module, and a PCF8574-driven 16x2 I2C LCD. I further confirm that all referenced technical datasheets, open-source C++ software libraries, and third-party engineering resources have been fully cited and acknowledged.")

    decl_p = doc.add_paragraph()
    decl_p.paragraph_format.space_before = Pt(14)
    decl_p.paragraph_format.line_spacing = 1.2
    decl_p.add_run("Candidate Signature: _________________________________\n\n").font.bold = True
    decl_p.add_run("Name: [Student Name]\nRoll Number: [Roll Number]\nDate: ________________________")

    # =========================================================================
    # FRONT MATTER: 3. ACKNOWLEDGEMENTS & ABSTRACT
    # =========================================================================
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("ACKNOWLEDGEMENTS")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)

    add_body(doc, "The conceptualization and successful realization of this engineering project would not have been possible without the invaluable guidance, resources, and encouragement provided by numerous individuals and academic institutions.")
    add_body(doc, "First and foremost, I express my deepest intellectual gratitude to my esteemed project supervisor, **[Project Guide Name]**, for their rigorous technical counsel, continuous mentorship, and patient reviews throughout both the virtual simulation and physical hardware prototyping phases. Their insistence on metrological precision and honest sensor labelling fundamentally enriched the engineering depth of this work.")
    add_body(doc, "I extend my sincere appreciation to **[HOD Name]**, Head of the Department of **[Department Name]**, and the dedicated faculty members and laboratory technical staff of **[College / Institution Name]** for furnishing the essential computational facilities, electronic benchtop instruments, and breadboard hardware necessary to conduct this study.")
    add_body(doc, "I am profoundly indebted to the open-source engineering and educational communities surrounding Arduino, Autodesk Tinkercad Circuits, and the Adafruit Industries software development teams. The accessible development tools, comprehensive peripheral libraries, and robust documentation provided by these organizations form the bedrock upon which modern student engineers cultivate hands-on instrumentation mastery.")
    add_body(doc, "Finally, I extend my heartfelt gratitude to my family, peers, and fellow engineering classmates whose constant encouragement, constructive discussions, and collaborative spirit provided unwavering motivation throughout the duration of this capstone endeavor.")

    # ABSTRACT
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("ABSTRACT")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)

    add_body(doc, "Continuous, real-time monitoring of indoor environmental conditions—specifically ambient temperature, relative atmospheric humidity, and airborne gaseous contaminants—is vital for ensuring human respiratory safety, occupational productivity, and building ventilation efficiency. This report details the comprehensive design, virtual simulation, embedded firmware architecture, physical hardware realization, and metrological validation of an autonomous, low-cost environmental monitoring station powered by an 8-bit ATmega328P microcontroller on the Arduino Uno platform.")
    add_body(doc, "The project adopts a dual-phase implementation strategy. In the virtual prototyping phase, an interactive circuit simulation was modeled within Autodesk Tinkercad Circuits, combining an analog TMP36 temperature transducer, a six-terminal heated gas sensor model, and a 16x2 alphanumeric liquid crystal display interfaced via an Adafruit I2C backpack. C++ firmware was formulated to linearize sensor voltages and categorize gas responses across a multi-tier qualitative threshold logic.")
    add_body(doc, "In the subsequent physical implementation phase, the system was physically fabricated on an 830-tie-point solderless breadboard. The thermal channel was upgraded to a digital DHT11 composite sensor, delivering simultaneous temperature (0–50°C) and relative humidity (20–90% RH) telemetry over a single-bus digital protocol. The gas channel was populated with a commercial MQ-135 metal oxide semiconductor breakout module utilizing Tin Dioxide (SnO2) chemiresistive surface adsorption. Visual feedback is rendered locally on a 16x2 LCD via a PCF8574 I2C adapter, supplemented by diagnostic serial telemetry streamed at 9600 baud.")
    add_body(doc, "Crucially, the report maintains strict metrological integrity: the MQ-135 is documented as a broad-spectrum, cross-sensitive qualitative detector, and its mapped 0–100 response scale is rigorously distinguished from certified Air Quality Index (AQI) or particulate matter (PM2.5) concentrations. Complete wiring topologies, full annotated source codes, empirical benchtop test records, a component catalog, a technical glossary, and formal references are presented in full detail.")

    p_kw = doc.add_paragraph()
    p_kw.paragraph_format.space_before = Pt(8)
    p_kw.paragraph_format.space_after = Pt(6)
    r_kwh = p_kw.add_run("Keywords: ")
    r_kwh.font.name = 'Times New Roman'
    r_kwh.font.size = Pt(11)
    r_kwh.font.bold = True
    r_kwt = p_kw.add_run("Arduino Uno R3, Environmental Monitoring, DHT11 Sensor, MQ-135 Gas Sensor, TMP36 Transducer, I2C Protocol, PCF8574, Tinkercad Simulation, Chemiresistance, Non-Blocking Firmware.")
    r_kwt.font.name = 'Times New Roman'
    r_kwt.font.size = Pt(11)
    r_kwt.font.italic = True

    # =========================================================================
    # FRONT MATTER: 4. TABLE OF CONTENTS (NO LIST OF FIGURES, NO LIST OF TABLES)
    # =========================================================================
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("DOCUMENT STRUCTURE AND TABLE OF CONTENTS")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)

    add_body(doc, "This capstone project report is formally organized into seventeen structured technical sections and reference documentation, thoroughly detailing every phase of the engineering lifecycle without unverified abbreviations or missing content:")

    toc_headers = ["Section", "Technical Title and Curricular Focus", "Nominal Span"]
    toc_widths = [Inches(1.1), Inches(4.3), Inches(1.1)]
    toc_rows = [
        ["Front Matter", "Certificate, Student Declaration, Acknowledgements, Technical Abstract", "Pages 1 – 3"],
        ["1", "Introduction and Project Overview: Industrial Context, Edge Philosophy, and Educational Value", "Pages 5 – 6"],
        ["2", "Problem Statement, Engineering Objectives, and Formal Scope Boundaries", "Pages 6 – 7"],
        ["3", "Theoretical Foundations: Embedded MCUs, Transducers, ADC Quantization, and I2C Buses", "Pages 8 – 9"],
        ["4", "System Architecture, Functional/Non-Functional Specs, Block Diagrams, and Data Pipelines", "Pages 10 – 11"],
        ["5", "Hardware Component Analysis: ATmega328P Core, HD44780 LCD, and Breadboard Interconnects", "Pages 11 – 12"],
        ["6", "Transducer Physics: TMP36 Silicon Bandgap, DHT11 Hygrometry, and MQ-135 SnO2 Chemiresistance", "Pages 13 – 15"],
        ["7", "Virtual Prototyping & EDA Simulation Workflow in Autodesk Tinkercad Circuits", "Pages 15 – 16"],
        ["8", "Comprehensive Circuit Schematic, Master Pin Interconnections, and Grounding Architecture", "Pages 17 – 18"],
        ["9", "Firmware Architecture, Cooperative State Scheduling, Conversion Algorithms, and Fault Traps", "Pages 19 – 20"],
        ["10", "Complete Virtual Firmware Implementation (Annotated Tinkercad C++ Source Code)", "Pages 21 – 22"],
        ["11", "Physical Hardware Implementation, Bill of Materials, Benchtop Assembly, and Power Budget", "Pages 22 – 24"],
        ["12", "Production Physical Firmware Implementation, Toolchain Setup, and Dual-Line Visualization", "Pages 24 – 26"],
        ["13", "Experimental Testing, Multi-Stage Calibration Protocols, and Empirical Benchtop Test Logs", "Pages 26 – 27"],
        ["14", "Results Discussion, Sensor Limitations, Cross-Sensitivity Analysis, and Laboratory Safety", "Pages 27 – 28"],
        ["15", "Engineering Component Specification, Project Timeline Schedule, and Future Roadmap", "Pages 28 – 29"],
        ["16", "Comprehensive Project Conclusions, Architectural Syntheses, and Final Reflections", "Pages 29 – 30"],
        ["17", "Technical Glossary, Formal References, Standards, and Verification Checklist", "Pages 30 – 31"]
    ]
    add_styled_table(doc, None, toc_headers, toc_rows, toc_widths)
    add_body(doc, "**Note:** Page references reflect document flow within Microsoft Word under standard Letter layout with 1-inch margins and 1.15 line spacing.")

    print("Front matter generated successfully.")
    return doc

print("build_document ready.")
