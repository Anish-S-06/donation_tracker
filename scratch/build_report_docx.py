import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D0D5DD", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_formatted_runs(paragraph, text, base_font_size=12, base_font_name='Times New Roman', is_caption=False, default_color=None):
    # Regex to capture bold-italic (***text***), bold (**text**), italic (*text*), inline code (`text`), math ($text$)
    # Order matters: *** first, then **, then *, `, $
    pattern = re.compile(r'(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*|`.*?`|\$.*?\$|\[.*?\]\(.*?\))')
    tokens = pattern.split(text)
    
    for token in tokens:
        if not token:
            continue
        
        run = paragraph.add_run()
        run.font.name = base_font_name
        run.font.size = Pt(base_font_size)
        if default_color:
            run.font.color.rgb = default_color
            
        if is_caption:
            run.font.italic = True
            
        if token.startswith('***') and token.endswith('***') and len(token) >= 6:
            run.text = token[3:-3]
            run.font.bold = True
            run.font.italic = True
        elif token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run.text = token[2:-2]
            run.font.bold = True
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run.text = token[1:-1]
            run.font.italic = True
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run.text = token[1:-1]
            run.font.name = 'Consolas'
            run.font.size = Pt(base_font_size * 0.9)
            run.font.color.rgb = RGBColor(0x8B, 0x1A, 0x1A) # subtle dark crimson for code tokens
        elif token.startswith('$') and token.endswith('$') and len(token) >= 2:
            # Simple clean math string representation
            math_text = token[1:-1].replace(r'\notin', '∉').replace(r'\text{', '').replace(r'}', '')
            run.text = math_text
            run.font.italic = True
        elif token.startswith('[') and '](' in token and token.endswith(')'):
            match = re.match(r'\[(.*?)\]\((.*?)\)', token)
            if match:
                link_text, link_url = match.groups()
                run.text = f"{link_text} ({link_url})" if "http" in link_url else link_text
                run.font.color.rgb = RGBColor(0x00, 0x33, 0x99)
                run.font.underline = True
            else:
                run.text = token
        else:
            run.text = token

def parse_markdown_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    doc = docx.Document()

    # 1. Page Setup & Margins (Academic: Left 1.25", Right 1.0", Top 1.0", Bottom 1.0")
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.0)
        
        # Add running header & footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Seva Sankalp: Technical Project Documentation | Dept. of Data Engineering")
        hrun.font.name = "Times New Roman"
        hrun.font.size = Pt(8.5)
        hrun.font.italic = True
        hrun.font.color.rgb = RGBColor(0x7A, 0x86, 0x9A)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("B.Tech V Sem Project Report — MVGR College of Engineering (A)")
        frun.font.name = "Times New Roman"
        frun.font.size = Pt(9.0)
        frun.font.color.rgb = RGBColor(0x7A, 0x86, 0x9A)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0x1F, 0x24, 0x2E)

    lines = content.split('\n')
    i = 0
    total_lines = len(lines)
    first_heading = True

    in_code_block = False
    code_block_lines = []

    in_table = False
    table_lines = []

    while i < total_lines:
        line = lines[i].rstrip('\r\n')

        # Check for code blocks
        if line.strip().startswith('```'):
            if in_code_block:
                # End of code block
                in_code_block = False
                code_text = '\n'.join(code_block_lines)
                
                # Render code block in a shaded callout table
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                tbl.autofit = False
                tbl.columns[0].width = Inches(6.25)
                
                cell = tbl.cell(0, 0)
                set_cell_background(cell, "F4F5F7")
                set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
                
                tcPr = cell._tc.get_or_add_tcPr()
                borders = parse_xml(
                    f'<w:tcBorders {nsdecls("w")}>'
                    f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/>'
                    f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="003366"/>'
                    f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/>'
                    f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/>'
                    f'</w:tcBorders>'
                )
                tcPr.append(borders)
                
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(2)
                cp.paragraph_format.space_after = Pt(2)
                cp.paragraph_format.line_spacing = 1.05
                crun = cp.add_run(code_text)
                crun.font.name = 'Consolas'
                crun.font.size = Pt(9.0)
                crun.font.color.rgb = RGBColor(0x24, 0x29, 0x2F)
                
                # Add a subtle blank paragraph after table
                spacer = doc.add_paragraph()
                spacer.paragraph_format.space_before = Pt(0)
                spacer.paragraph_format.space_after = Pt(6)
                
                code_block_lines = []
            else:
                in_code_block = True
                code_block_lines = []
            i += 1
            continue

        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue

        # Check for tables
        if line.strip().startswith('|') and line.strip().endswith('|'):
            table_lines.append(line)
            # Peek if next line is also a table line
            if i + 1 < total_lines and lines[i+1].strip().startswith('|') and lines[i+1].strip().endswith('|'):
                i += 1
                continue
            else:
                # Process the complete table
                parsed_rows = []
                for tline in table_lines:
                    # check if separator row like | :--- | :---: |
                    cleaned = re.sub(r'[\s|:\-]+', '', tline)
                    if cleaned == '':
                        continue # separator row
                    cells = [c.strip() for c in tline.strip().split('|')[1:-1]]
                    parsed_rows.append(cells)
                
                if parsed_rows:
                    num_rows = len(parsed_rows)
                    num_cols = max(len(r) for r in parsed_rows)
                    
                    tbl = doc.add_table(rows=num_rows, cols=num_cols)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    tbl.autofit = True
                    set_table_borders(tbl, color="D0D5DD", sz="4")
                    
                    for row_idx, rdata in enumerate(parsed_rows):
                        row = tbl.rows[row_idx]
                        is_header = (row_idx == 0)
                        
                        # Repeat header row across pages
                        if is_header:
                            trPr = row._tr.get_or_add_trPr()
                            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
                        
                        for col_idx in range(num_cols):
                            cell_val = rdata[col_idx] if col_idx < len(rdata) else ""
                            cell = row.cells[col_idx]
                            set_cell_margins(cell, top=100, bottom=100, left=130, right=130)
                            
                            if is_header:
                                set_cell_background(cell, "003366")
                            elif row_idx % 2 == 1:
                                set_cell_background(cell, "F9FAFB")
                            else:
                                set_cell_background(cell, "FFFFFF")
                                
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            p.paragraph_format.line_spacing = 1.05
                            
                            if is_header:
                                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                add_formatted_runs(p, cell_val, base_font_size=10.5, default_color=RGBColor(0xFF, 0xFF, 0xFF))
                                for r in p.runs:
                                    r.font.bold = True
                            else:
                                # Alignment: center if short or number
                                if len(cell_val) <= 6 or re.match(r'^\d+$', cell_val.strip()):
                                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                else:
                                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                                add_formatted_runs(p, cell_val, base_font_size=10.0)

                    spacer = doc.add_paragraph()
                    spacer.paragraph_format.space_before = Pt(0)
                    spacer.paragraph_format.space_after = Pt(6)

                table_lines = []
                i += 1
                continue

        # Horizontal rule -> Page break
        if line.strip() in ['---', '***', '___']:
            doc.add_page_break()
            i += 1
            continue

        # Blank line
        if not line.strip():
            i += 1
            continue

        # Heading 1 (# Heading)
        if line.startswith('# '):
            h_text = line[2:].strip()
            if not first_heading:
                doc.add_page_break()
            else:
                first_heading = False

            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.keep_with_next = True
            
            run = p.add_run(h_text.upper())
            run.font.name = 'Times New Roman'
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x00, 0x20, 0x60) # Deep Academic Navy
            
            i += 1
            continue

        # Heading 2 (## Heading)
        if line.startswith('## '):
            h_text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            
            run = p.add_run(h_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
            
            i += 1
            continue

        # Heading 3 (### Heading)
        if line.startswith('### '):
            h_text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            
            run = p.add_run(h_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
            
            i += 1
            continue

        # Figure / Table caption lines (*Figure 4.1:...*)
        if line.strip().startswith('*Figure ') or line.strip().startswith('*Table ') or (line.strip().startswith('*') and line.strip().endswith('*') and len(line.strip()) < 120 and 'Figure' in line):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(10)
            cap_text = line.strip().strip('*')
            run = p.add_run(cap_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x47, 0x54, 0x67)
            i += 1
            continue

        # Blockquote (> text)
        if line.startswith('> '):
            bq_text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.right_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            
            add_formatted_runs(p, bq_text, base_font_size=11, default_color=RGBColor(0x34, 0x40, 0x54))
            for r in p.runs:
                r.font.italic = True
            i += 1
            continue

        # Bullet List (* item or - item)
        bullet_match = re.match(r'^(\s*)[*\-]\s+(.*)', line)
        if bullet_match:
            indent_spaces, item_text = bullet_match.groups()
            indent_level = len(indent_spaces) // 2
            
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent = Inches(0.35 + indent_level * 0.25)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            
            add_formatted_runs(p, item_text, base_font_size=12)
            i += 1
            continue

        # Numbered List (1. item, 2. item)
        num_match = re.match(r'^(\s*)(\d+)\.\s+(.*)', line)
        if num_match:
            indent_spaces, num_str, item_text = num_match.groups()
            indent_level = len(indent_spaces) // 2
            
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4 + indent_level * 0.25)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            
            num_run = p.add_run(f"{num_str}. ")
            num_run.font.name = 'Times New Roman'
            num_run.font.size = Pt(12)
            num_run.font.bold = True
            
            add_formatted_runs(p, item_text, base_font_size=12)
            i += 1
            continue

        # Normal Paragraph
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        add_formatted_runs(p, line, base_font_size=12)
        i += 1

    # Save document
    doc.save(docx_path)
    print(f"Document successfully created at: {docx_path}")

if __name__ == '__main__':
    md_file = r"C:\Users\SANDAKA ANISH NIHAAL\.gemini\antigravity\brain\0c0befe5-ef6a-4fa5-bcf3-701edc32bbd2\Seva_Sankalp_Complete_Project_Report.md"
    out_docx = r"C:\Users\SANDAKA ANISH NIHAAL\donation_tracker\Seva_Sankalp_Complete_Project_Report.docx"
    brain_docx = r"C:\Users\SANDAKA ANISH NIHAAL\.gemini\antigravity\brain\0c0befe5-ef6a-4fa5-bcf3-701edc32bbd2\Seva_Sankalp_Complete_Project_Report.docx"
    
    parse_markdown_to_docx(md_file, out_docx)
    parse_markdown_to_docx(md_file, brain_docx)
